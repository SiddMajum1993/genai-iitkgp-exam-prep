#!/usr/bin/env python3
"""Extract text (and note where images live) from a .pptx file."""
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def iter_shape_text(shape, depth=0):
    """Recursively walk a shape and yield (kind, text) tuples."""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            yield from iter_shape_text(s, depth + 1)
        return
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            line = "".join(run.text for run in para.runs).strip()
            if line:
                yield ("text", line)
    if shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                yield ("table_row", " | ".join(cells))
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        yield ("image", f"[image {shape.width}x{shape.height}]")


def extract(path):
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"\n===== SLIDE {i} =====")
        # Slide layout hint
        try:
            out.append(f"[layout: {slide.slide_layout.name}]")
        except Exception:
            pass
        image_count = 0
        for shape in slide.shapes:
            for kind, txt in iter_shape_text(shape):
                if kind == "image":
                    image_count += 1
                    out.append(f"[IMG {image_count}]")
                else:
                    out.append(txt)
        # Notes
        if slide.has_notes_slide:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                out.append(f"\n-- SPEAKER NOTES --\n{note}")
        if image_count:
            out.append(f"\n[slide contains {image_count} image(s)]")
    return "\n".join(out)


if __name__ == "__main__":
    for path in sys.argv[1:]:
        text = extract(path)
        outpath = path.rsplit(".", 1)[0] + ".txt"
        with open(outpath, "w") as f:
            f.write(text)
        print(f"{path} -> {outpath} ({len(text)} chars)")
